import streamlit as st
from pathlib import Path
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import PyPDF2
from docx import Document
from openpyxl import Workbook
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

# -----------------------------
# Risk level tags
# -----------------------------
def similarity_risk_level(score: float) -> str:
    if score >= 85:
        return "HIGH RISK 🔴"
    elif score >= 70:
        return "MEDIUM RISK 🟠"
    elif score >= 50:
        return "MILD SIMILARITY 🟡"
    else:
        return "SAFE 🟢"

# Plagiarism badge based on highest similarity
def plagiarism_badge(score: float) -> str:
    if score >= 90:
        return "Plagiarized ⚠️"
    elif score >= 75:
        return "Highly Similar 🟠"
    elif score >= 50:
        return "Possibly Inspired 🟡"
    else:
        return "Likely Original ✅"

# -----------------------------
# Text extraction
# -----------------------------
def extract_text(file):
    name = file.name.lower()
    if name.endswith(".txt"):
        return file.read().decode("utf-8", errors="ignore")
    elif name.endswith(".pdf"):
        reader = PyPDF2.PdfReader(file)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text
    elif name.endswith(".docx"):
        doc = Document(file)
        return "\n".join(p.text for p in doc.paragraphs)
    return ""

# -----------------------------
# Similarity computation
# -----------------------------
def compute_similarity(texts, enhanced=False):
    ngram = (1, 5) if enhanced else (1, 3)
    vectorizer = TfidfVectorizer(
        lowercase=True,
        ngram_range=ngram,
        stop_words=None
    )
    tfidf = vectorizer.fit_transform(texts)
    return cosine_similarity(tfidf)

# -----------------------------
# Excel report generator
# -----------------------------
def create_excel_report(results):
    wb = Workbook()
    ws = wb.active
    ws.title = "Similarity Report"
    ws.append(["File 1", "File 2", "Similarity (%)", "Risk Level"])
    for f1, f2, score in results:
        ws.append([f1, f2, f"{score:.2f}%", similarity_risk_level(score)])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# -----------------------------
# PDF report generator
# -----------------------------
def create_pdf_report(results, highest, lowest, avg, num_docs, num_comparisons):
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    x_margin = 40
    y = height - 50

    c.setFont("Helvetica-Bold", 16)
    c.drawString(x_margin, y, "Document Similarity Report")
    y -= 30

    c.setFont("Helvetica", 11)
    c.drawString(x_margin, y, f"Total Documents   : {num_docs}")
    y -= 15
    c.drawString(x_margin, y, f"Total Comparisons : {num_comparisons}")
    y -= 15
    c.drawString(x_margin, y, f"Highest Similarity: {highest[0]} <-> {highest[1]} = {highest[2]:.2f}%")
    y -= 15
    c.drawString(x_margin, y, f"Lowest Similarity : {lowest[0]} <-> {lowest[1]} = {lowest[2]:.2f}%")
    y -= 15
    c.drawString(x_margin, y, f"Average Similarity: {avg:.2f}%")
    y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(x_margin, y, "Sorted Results:")
    y -= 20
    c.setFont("Helvetica", 10)

    # Only print first N rows to avoid huge PDFs
    max_rows = 80
    count = 0

    for f1, f2, score in results:
        line = f"{f1} <-> {f2} : {score:.2f}% [{similarity_risk_level(score)}]"
        if y < 60:
            c.showPage()
            y = height - 50
            c.setFont("Helvetica", 10)
        c.drawString(x_margin, y, line)
        y -= 12
        count += 1
        if count >= max_rows:
            c.drawString(x_margin, y, "... (truncated additional rows)")
            break

    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer

# -----------------------------
# Streamlit GUI
# -----------------------------

st.set_page_config(page_title="Document Similarity Checker", layout="wide")

# -----------------------------
# PROJECT BRANDING HEADER
# -----------------------------
st.markdown(
    """
    <div style="text-align:center; padding: 12px; border-radius: 10px;
                background-color:#f0f2f6; border: 1px solid #ccc;">
        <h1 style="margin-bottom:5px;">📄 Document Similarity Checker</h1>
        <h3 style="color:#333; margin:0;">By <b>Shailesh Paul</b></h3>
        <h4 style="color:#666; margin-top:5px;">B.Tech – AIML</h4>
    </div>
    <br>
    """,
    unsafe_allow_html=True
)




mode = st.radio(
    "Select comparison mode:",
    ["Standard (TF-IDF)", "Enhanced Semantic-like (Lightweight)"]
)

threshold = st.slider("Similarity Threshold (%)", 0, 100, 50)

uploaded_files = st.file_uploader(
    "Upload documents (TXT / PDF / DOCX)",
    type=["txt", "pdf", "docx"],
    accept_multiple_files=True
)

if st.button("Run Similarity Analysis"):
    if not uploaded_files or len(uploaded_files) < 2:
        st.warning("Please upload at least two documents.")
    else:
        texts = []
        names = []
        name_to_text = {}

        for f in uploaded_files:
            text = extract_text(f)
            if text.strip():
                texts.append(text)
                names.append(f.name)
                name_to_text[f.name] = text
            else:
                st.warning(f"⚠ {f.name} could not be processed (empty text).")

        if len(texts) < 2:
            st.warning("Not enough valid documents to compare.")
        else:
            enhanced = (mode == "Enhanced Semantic-like (Lightweight)")
            sim_matrix = compute_similarity(texts, enhanced=enhanced)

            results = []
            for i in range(len(names)):
                for j in range(i + 1, len(names)):
                    score = sim_matrix[i, j] * 100
                    results.append((names[i], names[j], score))

            results.sort(key=lambda x: x[2], reverse=True)

            highest = results[0]
            lowest = results[-1]
            avg = sum(r[2] for r in results) / len(results)
            num_docs = len(names)
            num_comparisons = len(results)

            st.markdown("---")
            st.header("📊 Summary")

            col1, col2, col3 = st.columns(3)
            col1.metric("Total Documents", num_docs)
            col2.metric("Total Comparisons", num_comparisons)
            col3.metric("Average Similarity", f"{avg:.2f}%")

            badge = plagiarism_badge(highest[2])
            st.success(
                f"### 🏆 Highest Similarity\n"
                f"**{highest[0]}** ↔ **{highest[1]}** = **{highest[2]:.2f}%**\n\n"
                f"Plagiarism Status: **{badge}**"
            )

            st.markdown("---")
            st.subheader("📁 Sorted Similarity Results")
            for f1, f2, score in results:
                risk = similarity_risk_level(score)
                if score >= threshold:
                    st.write(f"✅ **{f1}** ↔ **{f2}** : **{score:.2f}%**  _[{risk}]_")
                else:
                    st.write(f"➖ {f1} ↔ {f2} : {score:.2f}%  _[{risk}]_")

            st.markdown("---")
            st.subheader("📥 Download Reports")

            # Excel download
            excel_buffer = create_excel_report(results)
            st.download_button(
                label="📊 Download Excel Report (.xlsx)",
                data=excel_buffer,
                file_name="similarity_report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

            # PDF download
            pdf_buffer = create_pdf_report(results, highest, lowest, avg, num_docs, num_comparisons)
            st.download_button(
                label="📄 Download PDF Report",
                data=pdf_buffer,
                file_name="similarity_report.pdf",
                mime="application/pdf"
            )

            st.markdown("---")
            st.subheader("🧾 Side-by-Side Document Comparison")

            # Default to highest pair
            default_1 = highest[0]
            default_2 = highest[1]

            col_a, col_b = st.columns(2)
            with col_a:
                doc1 = st.selectbox("Select first document", names, index=names.index(default_1))
            with col_b:
                doc2 = st.selectbox("Select second document", names, index=names.index(default_2))

            if doc1 == doc2:
                st.warning("Please select two different documents for comparison.")
            else:
                c1, c2 = st.columns(2)
                with c1:
                    st.write(f"### {doc1}")
                    st.text_area("Text - Document 1", name_to_text[doc1], height=350)
                with c2:
                    st.write(f"### {doc2}")
                    st.text_area("Text - Document 2", name_to_text[doc2], height=350)

